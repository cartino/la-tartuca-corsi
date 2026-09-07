from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

from .models import CourseAnalysis, CourseSlot, GroupMatch, WordCourseSection
from .normalization import ITALIAN_DAYS, italian_full_date
from .word_reader import COURSE_URL_RE

DAY_RE = re.compile("|".join(map(re.escape, ITALIAN_DAYS)), re.IGNORECASE)
TIME_RE = re.compile(r"\b\d{1,2}\.\d{2}\b")
MONTHS_RE = "gennaio|febbraio|marzo|aprile|maggio|giugno|luglio|agosto|settembre|ottobre|novembre|dicembre"
FULL_DATE_RE = re.compile(rf"\b\d{{1,2}}\s+(?:{MONTHS_RE})\s+20\d{{2}}\b", re.IGNORECASE)
NUMERIC_DATE_RE = re.compile(r"\b\d{1,2}/\d{1,2}(?:/20\d{2})?\b")
AGE_RE = re.compile(r"\b\d+\s*[-–]\s*\d+\s*ANNI\b|\b\d+\s*ANNI\b", re.IGNORECASE)


@dataclass(frozen=True)
class Replacement:
    start: int
    end: int
    value: str
    highlight: bool


def _apply_replacements(text: str, replacements: list[Replacement]) -> tuple[str, list[tuple[int, int]]]:
    replacements = sorted(replacements, key=lambda item: item.start)
    output: list[str] = []
    highlights: list[tuple[int, int]] = []
    cursor = 0
    new_pos = 0
    for item in replacements:
        if item.start < cursor:
            continue
        unchanged = text[cursor:item.start]
        output.append(unchanged)
        new_pos += len(unchanged)
        output.append(item.value)
        if item.highlight:
            highlights.append((new_pos, new_pos + len(item.value)))
        new_pos += len(item.value)
        cursor = item.end
    output.append(text[cursor:])
    return "".join(output), highlights


def _replace_frequency(old: str, slot: CourseSlot) -> tuple[str, list[tuple[int, int]]]:
    day_match = DAY_RE.search(old)
    times = list(TIME_RE.finditer(old))
    if not day_match or len(times) < 2:
        new = f"Frequenza: settimanale, il {slot.day} dalle {slot.start_time} alle {slot.end_time}"
        return new, [(0, len(new))]

    replacements = [
        Replacement(day_match.start(), day_match.end(), slot.day, day_match.group(0).lower() != slot.day.lower()),
        Replacement(times[0].start(), times[0].end(), slot.start_time, times[0].group(0) != slot.start_time),
        Replacement(times[1].start(), times[1].end(), slot.end_time, times[1].group(0) != slot.end_time),
    ]
    return _apply_replacements(old, replacements)


def _replace_activation(old: str, slot: CourseSlot) -> tuple[str, list[tuple[int, int]]]:
    if not slot.activation_date:
        if slot.activation_raw:
            new = f"Attivazione: {slot.activation_raw}"
            return new, [(0, len(new))] if new != old else []
        return old, []

    day_match = DAY_RE.search(old)
    date_match = FULL_DATE_RE.search(old) or NUMERIC_DATE_RE.search(old)
    times = list(TIME_RE.finditer(old))
    full_date = italian_full_date(slot.activation_date)

    if not day_match or not date_match:
        new = f"Attivazione: {slot.day} {full_date} ore {slot.start_time}"
        return new, [(0, len(new))]

    replacements = [
        Replacement(day_match.start(), day_match.end(), slot.day, day_match.group(0).lower() != slot.day.lower()),
        Replacement(date_match.start(), date_match.end(), full_date, date_match.group(0).lower() != full_date.lower()),
    ]
    if times:
        last_time = times[-1]
        replacements.append(Replacement(last_time.start(), last_time.end(), slot.start_time, last_time.group(0) != slot.start_time))
    return _apply_replacements(old, replacements)


def _replace_age_label(old: str, new_age: str) -> tuple[str, list[tuple[int, int]]]:
    if not new_age:
        return old, []
    match = AGE_RE.search(old)
    new_display = new_age.upper()
    if not match:
        return old, []
    new = old[:match.start()] + new_display + old[match.end():]
    if new == old:
        return old, []
    return new, [(match.start(), match.start() + len(new_display))]


def _rewrite_paragraph(paragraph, new_text: str, highlight_spans: list[tuple[int, int]], bold_all: bool = False) -> None:
    old_text = paragraph.text
    old_colon = old_text.find(":")
    consumed = 0
    had_bold_label = False
    for existing_run in paragraph.runs:
        if old_colon >= 0 and consumed <= old_colon and existing_run.bold:
            had_bold_label = True
        consumed += len(existing_run.text)
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    label_end = new_text.find(":") + 1
    boundaries = {0, len(new_text)}
    if label_end > 0:
        boundaries.add(label_end)
    for start, end in highlight_spans:
        boundaries.add(max(0, min(len(new_text), start)))
        boundaries.add(max(0, min(len(new_text), end)))
    ordered = sorted(boundaries)
    for start, end in zip(ordered, ordered[1:]):
        if start == end:
            continue
        piece = new_text[start:end]
        run = paragraph.add_run(piece)
        if bold_all or (had_bold_label and end <= label_end):
            run.bold = True
        if any(start >= h_start and end <= h_end for h_start, h_end in highlight_spans):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _clear_paragraph(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def _set_highlighted_paragraph(paragraph, text: str, bold: bool = False) -> None:
    _clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _insert_before(anchor, text: str = "", highlighted: bool = False, bold: bool = False):
    p = anchor.insert_paragraph_before(text)
    if highlighted:
        _set_highlighted_paragraph(p, text, bold=bold)
    elif bold and p.runs:
        p.runs[0].bold = True
    return p


def _find_first_course_index(doc: Document) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        if COURSE_URL_RE.search(p.text.strip()):
            return i
    return None


def _strip_preamble(doc: Document) -> None:
    first = _find_first_course_index(doc)
    if first is None:
        return
    for p in list(doc.paragraphs[:first]):
        p._element.getparent().remove(p._element)


def _find_url_paragraph(doc: Document, url: str):
    for p in doc.paragraphs:
        if p.text.strip() == url.strip():
            return p
    return None


def _next_course_anchor(doc: Document, url: str):
    paragraphs = doc.paragraphs
    current = None
    for i, p in enumerate(paragraphs):
        if p.text.strip() == url.strip():
            current = i
            break
    if current is None:
        return None
    for p in paragraphs[current + 1:]:
        if COURSE_URL_RE.search(p.text.strip()):
            return p
    return None


def _insert_at_section_end(doc: Document, section: WordCourseSection, lines: list[tuple[str, bool, bool]]) -> None:
    anchor = _next_course_anchor(doc, section.url)
    if anchor is not None:
        for text, highlighted, bold in lines:
            _insert_before(anchor, text, highlighted=highlighted, bold=bold)
    else:
        for text, highlighted, bold in lines:
            p = doc.add_paragraph(text)
            if highlighted:
                _set_highlighted_paragraph(p, text, bold=bold)
            elif bold and p.runs:
                p.runs[0].bold = True


def _add_general_preamble(doc: Document, general_lines: list[str]) -> None:
    first_idx = _find_first_course_index(doc)
    if first_idx is None:
        return
    anchor = doc.paragraphs[first_idx]
    lines = ["PAGINA “TUTTI I CORSI”", "https://www.latartuca.it/corsi/"] + general_lines + [""]
    for text in lines:
        # Inserendo sempre prima dell'anchor originale, l'ordine naturale viene mantenuto.
        p = anchor.insert_paragraph_before(text)
        if text:
            _set_highlighted_paragraph(p, text, bold=text.startswith("PAGINA") or text.isupper())


def _default_group_label(slot: CourseSlot, idx: int) -> str:
    if slot.age_group:
        return f"GRUPPO {slot.age_group.upper()}"
    return f"GRUPPO {idx + 1}" if idx else ""


def _group_lines(slot: CourseSlot, data: dict[str, Any], idx: int = 0) -> list[tuple[str, bool, bool]]:
    label = (data.get("label") or _default_group_label(slot, idx)).strip()
    lines: list[tuple[str, bool, bool]] = []
    if data.get("sede"):
        lines.append((f"Sede: {data['sede'].strip()}", True, True))
    if label:
        lines.append((label, True, True))
    lines.append((f"Frequenza: settimanale, il {slot.day} dalle {slot.start_time} alle {slot.end_time}", True, False))
    if data.get("durata"):
        lines.append((f"Durata: {data['durata'].strip()}", True, False))
    if data.get("quota"):
        lines.append((f"Quota: {data['quota'].strip()}", True, False))
    if data.get("insegnante"):
        lines.append((f"Insegnante: {data['insegnante'].strip()}", True, False))
    if slot.activation_date:
        lines.append((f"Attivazione: {slot.day} {italian_full_date(slot.activation_date)} ore {slot.start_time}", True, False))
    elif slot.activation_raw:
        lines.append((f"Attivazione: {slot.activation_raw}", True, False))
    if data.get("calendario"):
        lines.append((f"Calendario: {data['calendario'].strip()}", True, False))
    return lines


def _course_needs_manual_mapping(analysis: CourseAnalysis) -> bool:
    return any(gm.needs_confirmation or gm.old_group_index is None or gm.new_slot_index is None for gm in analysis.group_matches)


def _resolved_group_matches(analysis: CourseAnalysis, section: WordCourseSection, decisions: dict[str, Any]) -> list[GroupMatch]:
    """Ricostruisce le associazioni effettive scelte nella maschera v0.3."""
    if not _course_needs_manual_mapping(analysis):
        return list(analysis.group_matches)
    sec_idx = analysis.word_section_index
    if sec_idx is None:
        return []
    result: list[GroupMatch] = []
    used_old: set[int] = set()
    for slot_idx in analysis.pdf_slot_indexes:
        value = decisions.get(f"mapping::{sec_idx}::{slot_idx}", "")
        if value == "NEW":
            result.append(GroupMatch(None, slot_idx, 0, "nuovo gruppo", False))
            continue
        if value in ("", "SPECIAL", None):
            continue
        try:
            old_idx = int(value)
        except Exception:
            continue
        used_old.add(old_idx)
        score, reason = 0, "associazione scelta dalla responsabile"
        for gm in analysis.group_matches:
            if gm.old_group_index == old_idx and gm.new_slot_index == slot_idx:
                score, reason = gm.score, gm.reason
                break
        result.append(GroupMatch(old_idx, slot_idx, score, reason, False))
    for old_idx in range(len(section.groups)):
        if old_idx not in used_old:
            result.append(GroupMatch(old_idx, None, 0, "gruppo precedente non associato", True))
    return result


def generate_updated_word_v03(
    reference_docx_bytes: bytes,
    slots: list[CourseSlot],
    sections: list[WordCourseSection],
    analyses: list[CourseAnalysis],
    decisions: dict[str, Any],
) -> bytes:
    """Genera il Word operativo v0.3.

    `decisions` contiene solo le scelte raccolte dall'interfaccia. Le modifiche certe
    (frequenza/attivazione dei gruppi associati) vengono applicate automaticamente.
    """
    doc = Document(io.BytesIO(reference_docx_bytes))
    paragraphs = doc.paragraphs

    # 1) Modifiche sui paragrafi esistenti, prima di inserire/eliminare blocchi.
    for analysis in analyses:
        if analysis.word_section_index is None or not analysis.pdf_slot_indexes:
            continue
        section = sections[analysis.word_section_index]
        for gm in _resolved_group_matches(analysis, section, decisions):
            if gm.old_group_index is None or gm.new_slot_index is None:
                continue
            old_group = section.groups[gm.old_group_index]
            slot = slots[gm.new_slot_index]

            if old_group.frequency_paragraph_index is not None:
                p = paragraphs[old_group.frequency_paragraph_index]
                new, spans = _replace_frequency(p.text, slot)
                if new != p.text:
                    _rewrite_paragraph(p, new, spans)
            if old_group.activation_paragraph_index is not None:
                p = paragraphs[old_group.activation_paragraph_index]
                new, spans = _replace_activation(p.text, slot)
                if new != p.text:
                    _rewrite_paragraph(p, new, spans)
            elif slot.activation_date or slot.activation_raw:
                # Aggiunta gestita più avanti come riga extra, senza alterare indici.
                pass
            if old_group.label_paragraph_index is not None and slot.age_group:
                p = paragraphs[old_group.label_paragraph_index]
                new, spans = _replace_age_label(p.text, slot.age_group)
                if new != p.text:
                    _rewrite_paragraph(p, new, spans, bold_all=True)

            duration_key = f"duration::{analysis.word_section_index}::{gm.old_group_index}"
            duration_override = (decisions.get(duration_key) or "").strip()
            if duration_override and old_group.duration_paragraph_index is not None:
                p = paragraphs[old_group.duration_paragraph_index]
                text = duration_override if duration_override.lower().startswith("durata") else f"Durata: {duration_override}"
                if text != p.text:
                    _rewrite_paragraph(p, text, [(0, len(text))])

            calendar_key = f"calendar::{analysis.word_section_index}::{gm.old_group_index}"
            calendar_override = (decisions.get(calendar_key) or "").strip()
            if calendar_override and old_group.calendar_paragraph_indexes:
                first_idx = old_group.calendar_paragraph_indexes[0]
                p = paragraphs[first_idx]
                text = calendar_override if calendar_override.lower().startswith("calendario") else f"Calendario: {calendar_override}"
                _rewrite_paragraph(p, text, [(0, len(text))])
                for extra_idx in old_group.calendar_paragraph_indexes[1:]:
                    _rewrite_paragraph(paragraphs[extra_idx], "", [])

    # 2) Elimina eventuale parte generale/refuso precedente e ricostruisce le istruzioni.
    _strip_preamble(doc)
    general_lines: list[str] = []

    # 3) Gestione corsi nuovi e corsi assenti.
    for analysis in analyses:
        if analysis.status == "nuovo corso":
            data = decisions.get(f"newcourse::{analysis.key}", {}) or {}
            section_name = (data.get("sezione") or "DA DEFINIRE").strip()
            general_lines.append(f"INSERIRE IN SEZIONE “{section_name.upper()}”: - {analysis.pdf_name.upper()}")
        elif analysis.status == "non presente nel PDF" and analysis.word_section_index is not None:
            action = decisions.get(f"oldcourse::{analysis.word_section_index}", "")
            if action == "non_parte":
                general_lines.append(f"TOGLIERE DALLA PAGINA “TUTTI I CORSI”: - {analysis.word_title.upper()}")
            elif action == "rinomina":
                target = decisions.get(f"oldcourse_target::{analysis.word_section_index}", "")
                if target:
                    general_lines.append(f"AGGIORNARE/RINOMINARE “{analysis.word_title.upper()}” COME “{target.upper()}”")

    # 4) Istruzioni / nuovi gruppi dentro schede esistenti.
    for analysis in analyses:
        if analysis.word_section_index is None:
            continue
        section = sections[analysis.word_section_index]
        additions: list[tuple[str, bool, bool]] = []
        for gm in _resolved_group_matches(analysis, section, decisions):
            if gm.old_group_index is not None and gm.new_slot_index is None:
                action = decisions.get(f"oldgroup::{analysis.word_section_index}::{gm.old_group_index}", "")
                if action == "non_parte":
                    label = section.groups[gm.old_group_index].label_text or section.groups[gm.old_group_index].frequency_text
                    additions.append((f"TOGLIERE IL GRUPPO: {label}", True, True))
                elif action == "rinomina":
                    target_idx = decisions.get(f"oldgroup_target::{analysis.word_section_index}::{gm.old_group_index}")
                    if target_idx is not None and str(target_idx) != "":
                        slot = slots[int(target_idx)]
                        label = section.groups[gm.old_group_index].label_text or "GRUPPO"
                        additions.append((f"AGGIORNARE {label} COME {slot.age_group or slot.source_title}", True, True))
                elif action == "confluito":
                    target_idx = decisions.get(f"oldgroup_target::{analysis.word_section_index}::{gm.old_group_index}")
                    if target_idx is not None and str(target_idx) != "":
                        slot = slots[int(target_idx)]
                        label = section.groups[gm.old_group_index].label_text or "GRUPPO"
                        additions.append((f"ACCORPARE/TOGLIERE {label}: confluito nel gruppo {slot.age_group or slot.source_title}", True, True))
                elif action == "altro":
                    note = (decisions.get(f"oldgroup_note::{analysis.word_section_index}::{gm.old_group_index}") or "").strip()
                    if note:
                        additions.append((note, True, True))
            elif gm.old_group_index is None and gm.new_slot_index is not None:
                data = decisions.get(f"newgroup::{analysis.word_section_index}::{gm.new_slot_index}", {}) or {}
                additions.extend(_group_lines(slots[gm.new_slot_index], data, idx=gm.new_slot_index))
            elif gm.old_group_index is not None and gm.new_slot_index is not None:
                old_group = section.groups[gm.old_group_index]
                if old_group.activation_paragraph_index is None:
                    slot = slots[gm.new_slot_index]
                    if slot.activation_date:
                        additions.append((f"Attivazione: {slot.day} {italian_full_date(slot.activation_date)} ore {slot.start_time}", True, False))
                    elif slot.activation_raw:
                        additions.append((f"Attivazione: {slot.activation_raw}", True, False))
        if additions:
            _insert_at_section_end(doc, section, [("", False, False)] + additions)

    # 5) Nuove schede complete in fondo al documento.
    for analysis in analyses:
        if analysis.status != "nuovo corso":
            continue
        data = decisions.get(f"newcourse::{analysis.key}", {}) or {}
        doc.add_paragraph("")
        p = doc.add_paragraph(data.get("url", "URL NUOVA PAGINA: DA CREARE"))
        _set_highlighted_paragraph(p, p.text)
        title = (data.get("titolo") or analysis.pdf_name).upper()
        p = doc.add_paragraph(title)
        _set_highlighted_paragraph(p, title, bold=True)
        for paragraph_text in [x.strip() for x in (data.get("descrizione") or "").split("\n") if x.strip()]:
            p = doc.add_paragraph(paragraph_text)
            _set_highlighted_paragraph(p, paragraph_text)
        if data.get("sede"):
            p = doc.add_paragraph(f"Sede: {data['sede'].strip()}")
            _set_highlighted_paragraph(p, p.text, bold=True)
        for idx, slot_idx in enumerate(analysis.pdf_slot_indexes):
            slot = slots[slot_idx]
            group_data = {
                "label": data.get("group_label", "") if len(analysis.pdf_slot_indexes) == 1 else "",
                "durata": data.get("durata", ""),
                "quota": data.get("quota", ""),
                "insegnante": data.get("insegnante", ""),
                "calendario": data.get("calendario", ""),
            }
            for text, highlighted, bold in _group_lines(slot, group_data, idx):
                p = doc.add_paragraph(text)
                if highlighted:
                    _set_highlighted_paragraph(p, text, bold=bold)

    # Inserisce la parte generale per ultima, ora che tutti i contenuti sono presenti.
    _add_general_preamble(doc, general_lines)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
