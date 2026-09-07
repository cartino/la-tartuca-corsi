from __future__ import annotations

import io
from collections import Counter
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .models import CourseAnalysis


def generate_control_report(analyses: list[CourseAnalysis], decisions: dict[str, Any], blocking_errors: list[str], warnings: list[str]) -> bytes:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("La Tartuca - Report di controllo")
    run.bold = True
    run.font.size = Pt(20)

    counts = Counter(a.status for a in analyses)
    doc.add_heading("Riepilogo", level=1)
    for label, key in [
        ("Corsi esistenti", "esistente"),
        ("Corsi esistenti con nuovi gruppi", "esistente + nuovo gruppo"),
        ("Corsi esistenti con gruppi assenti", "esistente + gruppo assente"),
        ("Corsi nuovi", "nuovo corso"),
        ("Schede non presenti nel PDF", "non presente nel PDF"),
    ]:
        doc.add_paragraph(f"{label}: {counts.get(key, 0)}")
    doc.add_paragraph(f"Errori bloccanti: {len(blocking_errors)}")
    doc.add_paragraph(f"Avvisi: {len(warnings)}")

    if blocking_errors:
        doc.add_heading("Errori bloccanti", level=1)
        for item in blocking_errors:
            doc.add_paragraph(item, style="List Bullet")
    if warnings:
        doc.add_heading("Avvisi", level=1)
        for item in warnings:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Decisioni registrate", level=1)
    relevant = [(k, v) for k, v in decisions.items() if v not in (None, "", False, {}, [])]
    if not relevant:
        doc.add_paragraph("Nessuna decisione manuale registrata.")
    else:
        for key, value in relevant:
            if key.startswith("mapping::"):
                label = "Sistemazione gruppi"
            elif key.startswith("newcourse::"):
                label = "Dati nuovo corso"
            elif key.startswith("newgroup::"):
                label = "Dati nuovo gruppo"
            elif key.startswith("oldcourse::"):
                label = "Gestione corso non presente"
            elif key.startswith("oldgroup::"):
                label = "Gestione gruppo precedente non associato"
            elif key.startswith("oldgroup_target::"):
                label = "Destinazione gruppo confluito/rinominato"
            elif key.startswith("oldgroup_note::"):
                label = "Nota caso particolare"
            elif key.startswith("duration::"):
                label = "Correzione durata"
            elif key.startswith("calendar::"):
                label = "Correzione calendario"
            else:
                continue
            doc.add_paragraph(f"{label}: {value}", style="List Bullet")

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
